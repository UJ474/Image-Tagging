"""
generate_report_and_presentation.py
Generates:
  reports/phase1_report.pdf   — professional multi-page PDF report
  reports/phase1_slides.pptx  — PowerPoint presentation (10 slides)

Run from the repository root with the .venv active:
  python scripts/generate_report_and_presentation.py
"""

import json, os, sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
RESULTS_DIR = ROOT / "reports" / "results"
OUT_DIR = ROOT / "reports"

# ─── Load all 8 result JSONs ────────────────────────────────────────────────
def load_results():
    keys = [
        ("logreg", "pixels"), ("logreg", "hog"),
        ("knn",    "pixels"), ("knn",    "hog"),
        ("rf",     "pixels"), ("rf",     "hog"),
        ("svm",    "pixels"), ("svm",    "hog"),
    ]
    results = {}
    for model, feat in keys:
        p = RESULTS_DIR / f"{model}_{feat}.json"
        with open(p) as f:
            results[(model, feat)] = json.load(f)
    return results

CLASS_NAMES = [
    "T-shirt/top", "Trouser", "Pullover", "Dress", "Coat",
    "Sandal", "Shirt", "Sneaker", "Bag", "Ankle boot"
]

MODEL_LABELS = {
    "logreg": "Logistic Regression",
    "knn":    "k-NN",
    "rf":     "Random Forest",
    "svm":    "Linear SVM",
}

# ═══════════════════════════════════════════════════════════════════════════
# PART 1 — PDF REPORT via ReportLab
# ═══════════════════════════════════════════════════════════════════════════
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm, mm
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    HRFlowable, PageBreak, KeepTogether
)

# ── Colour palette ──────────────────────────────────────────────────────────
C_NAVY    = colors.HexColor("#0f172a")
C_INDIGO  = colors.HexColor("#4f46e5")
C_SLATE   = colors.HexColor("#1e293b")
C_MID     = colors.HexColor("#475569")
C_LIGHT   = colors.HexColor("#94a3b8")
C_RULE    = colors.HexColor("#e2e8f0")
C_GREEN   = colors.HexColor("#059669")
C_RED     = colors.HexColor("#dc2626")
C_LBLUE   = colors.HexColor("#eff6ff")
C_CELL_HL = colors.HexColor("#e0e7ff")


def build_pdf(results: dict):
    out_path = OUT_DIR / "phase1_report.pdf"

    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=A4,
        leftMargin=2.5*cm, rightMargin=2.5*cm,
        topMargin=2.5*cm,  bottomMargin=2.5*cm,
        title="Phase 1 AML Baseline Report - Fashion-MNIST",
        author="Team T112",
    )

    W = A4[0] - 5*cm

    base = getSampleStyleSheet()

    def S(name, **kw):
        return ParagraphStyle(name, parent=base["Normal"], **kw)

    sTitle    = S("sTitle",    fontSize=26, leading=32, textColor=C_NAVY,
                  fontName="Helvetica-Bold", spaceAfter=4, alignment=TA_CENTER)
    sSubtitle = S("sSubtitle", fontSize=12, leading=16, textColor=C_MID,
                  fontName="Helvetica", alignment=TA_CENTER, spaceAfter=2)
    sMeta     = S("sMeta",     fontSize=9,  leading=13, textColor=C_LIGHT,
                  fontName="Helvetica", alignment=TA_CENTER, spaceAfter=4)
    sH1       = S("sH1",       fontSize=15, leading=19, textColor=C_INDIGO,
                  fontName="Helvetica-Bold", spaceBefore=18, spaceAfter=4)
    sH2       = S("sH2",       fontSize=12, leading=15, textColor=C_SLATE,
                  fontName="Helvetica-Bold", spaceBefore=12, spaceAfter=4)
    sH3       = S("sH3",       fontSize=10, leading=13, textColor=C_MID,
                  fontName="Helvetica-BoldOblique", spaceBefore=8, spaceAfter=3)
    sBody     = S("sBody",     fontSize=9.5, leading=14, textColor=C_SLATE,
                  fontName="Helvetica", alignment=TA_JUSTIFY, spaceAfter=6)
    sBullet   = S("sBullet",   fontSize=9.5, leading=14, textColor=C_SLATE,
                  fontName="Helvetica", leftIndent=14, spaceAfter=3)
    sCaption  = S("sCaption",  fontSize=8,  leading=11, textColor=C_LIGHT,
                  fontName="Helvetica-Oblique", alignment=TA_CENTER, spaceAfter=8)

    def HR():
        return HRFlowable(width="100%", thickness=0.5, color=C_RULE, spaceAfter=6, spaceBefore=4)

    def h1(txt): return Paragraph(txt, sH1)
    def h2(txt): return Paragraph(txt, sH2)
    def h3(txt): return Paragraph(txt, sH3)
    def p(txt):  return Paragraph(txt, sBody)
    def bullet(items):
        return [Paragraph(f"&#8226;&#160;&#160;{it}", sBullet) for it in items]

    def styled_table(data, col_widths=None, hl_rows=None, green_cells=None, red_cells=None):
        hl_rows     = hl_rows or []
        green_cells = green_cells or []
        red_cells   = red_cells   or []
        t = Table(data, colWidths=col_widths, repeatRows=1)
        style = [
            ("BACKGROUND",    (0,0), (-1,0), C_INDIGO),
            ("TEXTCOLOR",     (0,0), (-1,0), colors.white),
            ("FONTNAME",      (0,0), (-1,0), "Helvetica-Bold"),
            ("FONTSIZE",      (0,0), (-1,0), 8),
            ("BOTTOMPADDING", (0,0), (-1,0), 7),
            ("TOPPADDING",    (0,0), (-1,0), 7),
            ("FONTNAME",      (0,1), (-1,-1), "Helvetica"),
            ("FONTSIZE",      (0,1), (-1,-1), 8.5),
            ("TOPPADDING",    (0,1), (-1,-1), 5),
            ("BOTTOMPADDING", (0,1), (-1,-1), 5),
            ("ROWBACKGROUNDS",(0,1), (-1,-1), [colors.white, C_LBLUE]),
            ("LINEBELOW",     (0,0), (-1,-1), 0.4, C_RULE),
            ("ALIGN",         (0,0), (-1,-1), "CENTER"),
            ("ALIGN",         (0,0), (0,-1),  "LEFT"),
            ("VALIGN",        (0,0), (-1,-1), "MIDDLE"),
            ("BOX",           (0,0), (-1,-1), 0.5, C_RULE),
        ]
        for row in hl_rows:
            style += [("BACKGROUND", (0,row), (-1,row), C_CELL_HL),
                      ("TEXTCOLOR",  (0,row), (-1,row), C_INDIGO),
                      ("FONTNAME",   (0,row), (-1,row), "Helvetica-Bold")]
        for (r,c) in green_cells:
            style += [("TEXTCOLOR",  (c,r),(c,r), C_GREEN),
                      ("FONTNAME",   (c,r),(c,r), "Helvetica-Bold")]
        for (r,c) in red_cells:
            style.append(("TEXTCOLOR", (c,r),(c,r), C_RED))
        t.setStyle(TableStyle(style))
        return t

    story = []

    # ── Title page ──────────────────────────────────────────────────────
    story.append(Spacer(1, 1.5*cm))
    story.append(Paragraph("Phase 1 Report", sSubtitle))
    story.append(Paragraph("Applied Machine Learning Baseline", sTitle))
    story.append(Paragraph("Automatic Clothing Category Recognition on Fashion-MNIST", sSubtitle))
    story.append(Spacer(1, 0.5*cm))
    story.append(HRFlowable(width="60%", thickness=1.5, color=C_INDIGO, hAlign="CENTER", spaceAfter=10))
    story.append(Paragraph("Course Project 11  |  Team T112  |  Section D  |  October 2026", sMeta))
    story.append(Spacer(1, 0.4*cm))
    story.append(styled_table(
        [["#", "Member", "GitHub", "Contribution"],
         ["1", "Anant Jain",    "anantjain-afk",  "Repo setup, data pipeline, master notebook"],
         ["2", "Kartik Mehra",  "KartikMehra22",  "Feature extraction (Pixels & HOG), caching"],
         ["3", "Krishna Gehlot","NSTKrishna",      "Logistic Regression baseline"],
         ["4", "Vipul Sharma",  "vs-code06",       "k-NN baseline & latency profiling"],
         ["5", "Aditya",        "AdityaInnovates", "Random Forest baseline"],
         ["6", "Utkarsh Jain",  "UJ474",           "Linear SVM & garment error analysis"]],
        col_widths=[0.5*cm, 3*cm, 3.2*cm, W-6.7*cm]))
    story.append(PageBreak())

    # ── Section 1: Problem Statement ───────────────────────────────────
    story.append(h1("1. Problem Statement"))
    story.append(HR())
    story.append(p(
        "Commercial inventory systems must reliably tag thousands of clothing images per day. "
        "Fashion-MNIST provides a standardized 10-class benchmark (70,000 grayscale 28x28 images) "
        "representative of this task. Phase 1 evaluates four classical ML algorithms on two visual "
        "representations to establish a rigorous baseline and identify the exact failure modes that "
        "motivate a Phase 2 Convolutional Neural Network."))
    story.append(p(
        "<b>Primary research question:</b> Can fixed, handcrafted visual descriptors adequately "
        "classify clothing for a commercial inventory catalogue, or does fine-grained category "
        "confusion necessitate hierarchical visual representations learned by a CNN?"))

    # ── Section 2: Experimental Setup ──────────────────────────────────
    story.append(h1("2. Experimental Setup"))
    story.append(HR())
    story.append(h2("2.1  Dataset and Data Splits"))
    story.append(styled_table(
        [["Split", "Size", "Purpose"],
         ["Training (dev)", "48,000", "Model selection & hyperparameter tuning"],
         ["Validation",     "12,000", "Held-out evaluation during tuning"],
         ["Test (official)","10,000", "Final one-time evaluation (held-out)"]],
        col_widths=[3.5*cm, 2.5*cm, W-6*cm]))
    story.append(p("All splits use stratified sampling with random_state=42. Pixel values are "
                   "normalized from uint8 [0, 255] to float32 [0.0, 1.0]."))

    story.append(h2("2.2  Class Labels"))
    story.append(styled_table(
        [["ID","Class","ID","Class"],
         ["0","T-shirt/top","5","Sandal"],
         ["1","Trouser","6","Shirt (most challenging)"],
         ["2","Pullover","7","Sneaker"],
         ["3","Dress","8","Bag"],
         ["4","Coat","9","Ankle boot"]],
        col_widths=[1*cm, 3.5*cm, 1*cm, 3.5*cm],
        red_cells=[(2,1)]))

    story.append(h2("2.3  Feature Representations"))
    story.append(styled_table(
        [["Feature","Dimensions","Description"],
         ["Pixels","784-d","Flattened 28x28 normalized pixel values in [0.0, 1.0]"],
         ["HOG","324-d","9 orientations, 7x7 px/cell, 2x2 cells/block, L2-Hys normalization"]],
        col_widths=[2*cm, 2.2*cm, W-4.2*cm]))

    story.append(h2("2.4  Models and Best Hyperparameters"))
    story.append(styled_table(
        [["Model","Best Config (Pixels)","Best Config (HOG)"],
         ["Logistic Regression","C=0.1, solver=lbfgs","C=0.1, solver=lbfgs"],
         ["k-NN","k=7, weights=distance","k=5, weights=distance"],
         ["Random Forest","200 trees, max_depth=20","200 trees, max_depth=None"],
         ["Linear SVM","C=0.1, max_iter=2000","C=1.0, max_iter=2000"]],
        col_widths=[3.5*cm, (W-3.5*cm)/2, (W-3.5*cm)/2]))

    story.append(h2("2.5  Evaluation Metrics"))
    story += bullet([
        "<b>Accuracy</b> -- fraction of correct predictions on the 10,000 test images.",
        "<b>Macro F1</b> -- unweighted mean F1 across all 10 classes.",
        "<b>Per-class Recall</b> -- sensitivity per garment category.",
        "<b>Parameter Count</b> -- weights for linear models, MiB for k-NN, nodes for RF.",
        "<b>Inference Latency</b> -- median of 5 warmed-up calls on 1,000-image batches (ms/image).",
    ])
    story.append(PageBreak())

    # ── Section 3: Results ─────────────────────────────────────────────
    story.append(h1("3. Results"))
    story.append(HR())
    story.append(h2("3.1  Master Comparison Table"))

    sorted_rows = sorted(results.items(), key=lambda x: -x[1]["test_accuracy"])
    tbl_data = [["#","Model","Features","Test Acc","Macro F1","Parameters","Latency (ms)"]]
    green_c = []
    for i, ((m, f), r) in enumerate(sorted_rows, 1):
        params = r["parameter_count"]
        ps = (f"{params:.1f} MiB" if m=="knn" else
              (f"{int(params):,} nodes" if m=="rf" else f"{int(params):,}"))
        row_idx = i
        if i == 1:
            green_c += [(row_idx, 3), (row_idx, 4)]
        tbl_data.append([str(i), MODEL_LABELS[m], "Pixels" if f=="pixels" else "HOG",
                         f"{r['test_accuracy']*100:.2f}%",
                         f"{r['test_macro_f1']*100:.2f}%",
                         ps,
                         f"{r['inference_latency_ms_per_image']:.4f}"])

    story.append(styled_table(tbl_data,
        col_widths=[0.5*cm, 3.8*cm, 1.5*cm, 1.6*cm, 1.6*cm, 3*cm, 2*cm],
        hl_rows=[1], green_cells=green_c))
    story.append(Paragraph(
        "Table 1. All 8 model configurations ranked by test accuracy. "
        "Best configuration highlighted (Random Forest + Pixels, 87.43%).", sCaption))

    story.append(h2("3.2  Key Findings"))
    for f_text in [
        "<b>Finding 1: Random Forest dominates overall.</b>  RF on pixels achieves 87.43% accuracy, "
        "outperforming k-NN by +2.02 pp, Logistic Regression by +3.11 pp, and Linear SVM by +3.63 pp.",
        "<b>Finding 2: HOG helps linear models, hurts non-linear ones.</b>  For LogReg and SVM, HOG "
        "improves accuracy by +0.23-0.43 pp while reducing parameter count by 58.6% (7,850 -> 3,250). "
        "For RF and k-NN, HOG reduces accuracy by 3.08 pp and 1.57 pp.",
        "<b>Finding 3: k-NN memory and latency trade-off.</b>  Switching from pixels (179.4 MiB) to HOG "
        "(74.2 MiB) reduces k-NN memory by 58.7% and inference latency by 2.4x at only 1.57 pp accuracy cost.",
        "<b>Finding 4: Linear classifiers are fastest.</b>  Logistic Regression on HOG runs at "
        "0.0004 ms/image, 685x faster than k-NN on pixels.",
        "<b>Finding 5: Shirt (class 6) is universally hardest.</b>  Mean Shirt recall across all 8 "
        "configurations is 57.7%. No model exceeds 62.6%.",
    ]:
        story.append(p(f_text))
    story.append(PageBreak())

    # ── Section 4: Per-Class Recall ────────────────────────────────────
    story.append(h1("4. Per-Class Recall Analysis"))
    story.append(HR())
    story.append(h2("4.1  Best Model -- Random Forest (Pixels)"))

    r_best = results[("rf", "pixels")]
    rc_data = [["Class", "Label", "Recall", "Assessment"]]
    red_rc = []
    for idx, (cl, rec) in enumerate(zip(CLASS_NAMES, r_best["per_class_recall"])):
        asmt = ("Excellent" if rec >= 0.93 else
                "Good" if rec >= 0.80 else
                "Fair" if rec >= 0.70 else "Poor")
        row_i = idx + 1
        if asmt == "Poor":
            red_rc += [(row_i, 2), (row_i, 3)]
        rc_data.append([str(idx), cl, f"{rec*100:.1f}%", asmt])
    story.append(styled_table(rc_data,
        col_widths=[1*cm, 3.5*cm, 2*cm, W-6.5*cm], red_cells=red_rc))
    story.append(Paragraph(
        "Table 2. Per-class recall for the best model (Random Forest, Pixels, 87.43% accuracy).", sCaption))

    story.append(h2("4.2  Cross-Model Shirt Recall"))
    shirt_data = [["Model","Pixels","HOG"]]
    for m_ in ["logreg","knn","rf","svm"]:
        rp = results[(m_,"pixels")]["per_class_recall"][6]
        rh = results[(m_,"hog")]["per_class_recall"][6]
        shirt_data.append([MODEL_LABELS[m_], f"{rp*100:.1f}%", f"{rh*100:.1f}%"])
    mp = sum(results[(m_,"pixels")]["per_class_recall"][6] for m_ in ["logreg","knn","rf","svm"])/4
    mh = sum(results[(m_,"hog")]["per_class_recall"][6] for m_ in ["logreg","knn","rf","svm"])/4
    shirt_data.append(["Mean", f"{mp*100:.1f}%", f"{mh*100:.1f}%"])
    story.append(styled_table(shirt_data, col_widths=[4*cm, 3*cm, 3*cm],
        red_cells=[(i,1) for i in range(1,6)] + [(i,2) for i in range(1,6)]))
    story.append(Paragraph(
        "Table 3. Shirt (class 6) recall across all 8 configurations. No model exceeds 62.6%.", sCaption))
    story.append(PageBreak())

    # ── Section 5: Error Analysis ──────────────────────────────────────
    story.append(h1("5. Error Analysis: The Shirt Dilemma"))
    story.append(HR())
    story.append(h2("5.1  Confusion Pattern"))
    story.append(p(
        "Shirt (class 6) is the most visually ambiguous upper-body label. Its 28x28 silhouette "
        "overlaps with T-shirts (similar collar and sleeve length), Pullovers (similar torso "
        "coverage), and Coats (similar outer garment outline). The confusion is directional:"))
    story.append(styled_table(
        [["Error Pattern","Frequency"],
         ["Shirt predicted as T-shirt/top","~15-16% of Shirt test images"],
         ["Shirt predicted as Pullover","~12-14% of Shirt test images"],
         ["Pullover predicted as Coat","~8-10% of Pullover test images"]],
        col_widths=[8*cm, W-8*cm]))

    story.append(h2("5.2  Why Fixed Representations Fail"))
    story.append(h3("Flattened Pixels (784-d)"))
    story.append(p(
        "Linear classifiers assign one weight per pixel coordinate, with no notion that a collar "
        "shifted upward is the same discriminative cue. Small translations, pose variation, and "
        "contrast changes alter many coordinates while the garment category is unchanged."))
    story.append(h3("HOG Descriptors (324-d)"))
    story.append(p(
        "HOG improves on pixels by summarizing local edge orientations. However, it cannot learn "
        "that a button-row edge pattern near the neckline is more diagnostic than a similar edge "
        "elsewhere. Cell pooling removes exactly the fine spatial cues needed for upper-body "
        "distinction."))
    story.append(h3("Linear Decision Boundaries (LogReg, SVM)"))
    story.append(p(
        "A hyperplane cannot create discriminative evidence absent in the input representation. "
        "If Shirt and T-shirt/top share overlapping pixel or HOG distributions, changing the "
        "margin penalty C cannot resolve the overlap -- this explains the consistent recall "
        "plateau across all hyperparameter searches."))

    story.append(h2("5.3  Visual Evidence"))
    story += bullet([
        "svm_pixels_confusion_matrix.png -- Full 10-class normalized confusion matrix.",
        "svm_pixels_upper_body_confusion.png -- Focused 4x4 sub-matrix: T-shirt/Pullover/Coat/Shirt.",
        "svm_pixels_2_as_4_gallery.png -- Gallery: Pullover images misclassified as Coat.",
        "svm_pixels_6_as_0_gallery.png -- Gallery: Shirt images misclassified as T-shirt/top.",
    ])
    story.append(PageBreak())

    # ── Section 6: Efficiency ──────────────────────────────────────────
    story.append(h1("6. Efficiency Analysis"))
    story.append(HR())
    story.append(h2("6.1  Accuracy vs. Inference Latency"))
    lat_data = [["Model","Features","Latency (ms/img)","Test Accuracy"]]
    for (m,f), r in sorted(results.items(), key=lambda x: x[1]["inference_latency_ms_per_image"]):
        lat_data.append([MODEL_LABELS[m], "HOG" if f=="hog" else "Pixels",
                         f"{r['inference_latency_ms_per_image']:.4f}",
                         f"{r['test_accuracy']*100:.2f}%"])
    story.append(styled_table(lat_data, col_widths=[3.5*cm, 2*cm, 3.5*cm, 3.5*cm]))
    story.append(Paragraph(
        "Table 4. All configurations sorted by inference latency. "
        "Linear models are 100-700x faster than k-NN.", sCaption))

    story.append(h2("6.2  Deployment Guidance"))
    story.append(p("<b>Maximum accuracy:</b>  Random Forest on pixels -- 87.43% at 0.026 ms/image."))
    story.append(p("<b>Real-time throughput (&lt;1 ms budget):</b>  Logistic Regression on HOG -- "
                   "84.55% at 0.0004 ms/image, with only 3,250 parameters."))
    story.append(p("<b>Edge deployment:</b>  Logistic Regression or SVM on HOG -- 3,250 parameters "
                   "(vs. 7,850 for pixel variants)."))
    story.append(p("<b>k-NN caveat:</b>  Despite 85.41% accuracy on pixels, k-NN stores 179 MiB and "
                   "is 685x slower than LogReg. Not suitable for production scale."))
    story.append(PageBreak())

    # ── Section 7: Phase 2 Bridge ──────────────────────────────────────
    story.append(h1("7. Bridge to Phase 2"))
    story.append(HR())
    story.append(p(
        "The consistent Shirt confusion (mean recall 57.7%) across all 8 Phase 1 configurations "
        "is a representation problem, not a hyperparameter problem. A CNN directly addresses "
        "each identified limitation:"))
    story.append(styled_table(
        [["CNN Property","Phase 1 Limitation Addressed"],
         ["Convolutions share detectors across locations",
          "Collar/button-row edges are position-sensitive in pixel/HOG space"],
         ["Stacked layers compose features hierarchically",
          "Fixed HOG cannot integrate edges into semantic garment parts"],
         ["Pooling provides modest translation tolerance",
          "Pixel flattening is corrupted by small garment position shifts"],
         ["End-to-end learned features",
          "HOG vocabulary is fixed before seeing the task; cannot specialize"]],
        col_widths=[6*cm, W-6*cm]))

    story.append(h2("7.1  Phase 2 Success Criteria"))
    story += bullet([
        "Shirt recall above <b>62.6%</b> (the Phase 1 ceiling across all 8 models).",
        "Reduction in Shirt -> T-shirt/top and Shirt -> Pullover confusion entries "
        "in the focused 4x4 upper-body sub-matrix.",
        "Overall accuracy improvement over the best AML model (87.43%).",
        "Same held-out test split and metric suite for direct comparability.",
    ])

    # ── Section 8: Conclusion ──────────────────────────────────────────
    story.append(h1("8. Conclusion"))
    story.append(HR())
    story.append(styled_table(
        [["Dimension","Winner","Value"],
         ["Best Accuracy",     "Random Forest + Pixels",   "87.43%"],
         ["Best Macro F1",     "Random Forest + Pixels",   "87.31%"],
         ["Fastest Inference", "Logistic Regression + HOG","0.0004 ms/img"],
         ["Most Compact",      "LogReg / SVM + HOG",       "3,250 parameters"],
         ["Hardest Category",  "Shirt (class 6)",          "57.7% mean recall"]],
        col_widths=[4.5*cm, 5*cm, W-9.5*cm],
        hl_rows=[5], red_cells=[(5,2)]))
    story.append(Paragraph("Table 5. Phase 1 summary results.", sCaption))
    story.append(p(
        "Phase 1 establishes a rigorous, fully reproducible 8-configuration AML baseline on "
        "Fashion-MNIST. All four models plateau at the same fundamental limit: fine-grained "
        "upper-body garment categories require spatial feature composition that fixed descriptors "
        "and linear decision boundaries cannot provide. Phase 2 (CNN) will directly test whether "
        "learned hierarchical spatial features close the gap that eight Phase 1 configurations "
        "could not."))
    story.append(Spacer(1, 0.5*cm))
    story.append(HRFlowable(width="100%", thickness=0.5, color=C_RULE, spaceAfter=6))
    story.append(Paragraph(
        "All metrics evaluated on the official Fashion-MNIST held-out test set (10,000 images). "
        "Source data: reports/results/*.json. "
        "Full visualizations: notebooks/00_master_aml_comparison.ipynb.", sCaption))

    doc.build(story)
    print(f"[PDF] Written -> {out_path}")
    return out_path


# ═══════════════════════════════════════════════════════════════════════════
# PART 2 — PPTX via python-pptx
# ═══════════════════════════════════════════════════════════════════════════
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

PC_NAVY   = RGBColor(0x0f, 0x17, 0x2a)
PC_INDIGO = RGBColor(0x4f, 0x46, 0xe5)
PC_WHITE  = RGBColor(0xff, 0xff, 0xff)
PC_LIGHT  = RGBColor(0x94, 0xa3, 0xb8)
PC_GREEN  = RGBColor(0x05, 0x96, 0x69)
PC_RED    = RGBColor(0xdc, 0x26, 0x26)
PC_SLATE  = RGBColor(0x1e, 0x29, 0x3b)
PC_MID    = RGBColor(0x47, 0x55, 0x69)
PC_CELL_H = RGBColor(0xee, 0xf2, 0xff)
PC_ACCENT = RGBColor(0xa5, 0xb4, 0xfc)

W_IN = 13.33
H_IN = 7.5


def add_slide(prs):
    layout = prs.slide_layouts[6]  # blank
    return prs.slides.add_slide(layout)


def set_bg(slide, rgb):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = rgb


def add_rect(slide, x, y, w, h, fill_rgb, line_rgb=None):
    shp = slide.shapes.add_shape(1,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill_rgb
    if line_rgb:
        shp.line.color.rgb = line_rgb
        shp.line.width = Pt(0.5)
    else:
        shp.line.fill.background()
    return shp


def tb(slide, x, y, w, h):
    txb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = txb.text_frame
    tf.word_wrap = True
    return tf


def add_text(slide, x, y, w, h, lines, size=11, color=PC_LIGHT,
             bold=False, italic=False, align=PP_ALIGN.LEFT):
    tf = tb(slide, x, y, w, h)
    first = True
    for line in lines:
        p_ = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p_.alignment = align
        run = p_.add_run()
        run.text = line
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color


def pptx_table(slide, x, y, w, h, data, fs=10):
    rows = len(data)
    cols = len(data[0])
    t = slide.shapes.add_table(rows, cols,
        Inches(x), Inches(y), Inches(w), Inches(h)).table
    col_w = int(Inches(w) / cols)
    for c in range(cols):
        t.columns[c].width = col_w
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.text = ""
            tf_ = cell.text_frame
            tf_.paragraphs[0].alignment = PP_ALIGN.CENTER
            run = tf_.paragraphs[0].add_run()
            run.text = str(val)
            run.font.name = "Calibri"
            run.font.size = Pt(fs)
            if r == 0:
                run.font.bold = True
                run.font.color.rgb = PC_WHITE
                cell.fill.solid()
                cell.fill.fore_color.rgb = PC_INDIGO
            else:
                run.font.color.rgb = PC_SLATE
                cell.fill.solid()
                cell.fill.fore_color.rgb = PC_CELL_H if r % 2 == 0 else PC_WHITE
    return t


def slide_title(slide, text, y=0.6, size=26):
    add_text(slide, 0.5, y, W_IN - 0.8, 1.1, [text],
             size=size, color=PC_WHITE, bold=True)


def slide_tag(slide, text):
    add_text(slide, 0.5, 0.22, W_IN - 0.8, 0.3, [text.upper()],
             size=9, color=PC_ACCENT, bold=True)


def build_pptx(results: dict):
    out_path = OUT_DIR / "phase1_slides.pptx"
    prs = Presentation()
    prs.slide_width  = Inches(W_IN)
    prs.slide_height = Inches(H_IN)

    # ── Slide 1: Title ───────────────────────────────────────────────────
    sl = add_slide(prs)
    set_bg(sl, PC_NAVY)
    add_rect(sl, 0, 0, 0.18, H_IN, PC_INDIGO)
    slide_tag(sl, "Phase 1  |  AML Baseline  |  Team T112  |  Section D")
    slide_title(sl, "Automatic Clothing\nCategory Recognition", y=1.1, size=36)
    add_text(sl, 0.5, 3.1, 9, 0.5,
             ["Four classical ML algorithms x two visual representations on Fashion-MNIST"],
             size=14, color=PC_LIGHT)
    add_text(sl, 0.5, 3.9, 12.5, 3, [
        "1. Anant Jain (anantjain-afk)  --  Repo setup, data pipeline, master notebook",
        "2. Kartik Mehra (KartikMehra22)  --  Feature extraction (Pixels & HOG), caching",
        "3. Krishna Gehlot (NSTKrishna)  --  Logistic Regression baseline",
        "4. Vipul Sharma (vs-code06)  --  k-NN baseline & latency profiling",
        "5. Aditya (AdityaInnovates)  --  Random Forest baseline",
        "6. Utkarsh Jain (UJ474)  --  Linear SVM & garment error analysis",
    ], size=11, color=RGBColor(0x64, 0x74, 0x8b))

    # ── Slide 2: Problem Statement ────────────────────────────────────────
    sl = add_slide(prs)
    set_bg(sl, PC_NAVY)
    add_rect(sl, 0, 0, 0.18, H_IN, PC_INDIGO)
    slide_tag(sl, "Problem Statement")
    slide_title(sl, "Can handcrafted features classify clothing?")
    add_text(sl, 0.5, 1.25, 12, 0.3,
             ["Fashion-MNIST as a proxy for commercial inventory tagging"],
             size=13, color=PC_LIGHT)

    add_rect(sl, 0.5, 1.65, 5.9, 2.2, PC_SLATE)
    add_text(sl, 0.65, 1.75, 5.5, 0.3, ["THE DATASET"], size=8, color=PC_ACCENT, bold=True)
    add_text(sl, 0.65, 2.0, 5.5, 1.7, [
        "70,000 grayscale 28x28 images   |   10 clothing categories",
        "Train: 48,000  |  Val: 12,000  |  Test: 10,000",
        "Pixels normalized: float32 in [0.0, 1.0]",
        "Stratified split  (random_state=42)",
    ], size=11, color=PC_LIGHT)

    add_rect(sl, 0.5, 4.1, 5.9, 3.0, PC_SLATE)
    add_text(sl, 0.65, 4.2, 5.5, 0.3, ["RESEARCH QUESTION"], size=8, color=PC_ACCENT, bold=True)
    add_text(sl, 0.65, 4.55, 5.5, 2.4,
             ['"Can fixed, handcrafted visual descriptors adequately classify clothing, '
              'or does fine-grained category confusion necessitate hierarchical features '
              'learned by a CNN?"'],
             size=11, color=PC_LIGHT, italic=True)

    add_rect(sl, 6.7, 1.65, 6.3, 5.45, PC_SLATE)
    add_text(sl, 6.85, 1.75, 5.9, 0.3, ["CLASS LABELS (10)"], size=8, color=PC_ACCENT, bold=True)
    add_text(sl, 6.85, 2.05, 5.9, 5.0, [
        "0 -- T-shirt/top         5 -- Sandal",
        "1 -- Trouser             6 -- Shirt  (hardest)",
        "2 -- Pullover            7 -- Sneaker",
        "3 -- Dress               8 -- Bag",
        "4 -- Coat                9 -- Ankle boot",
    ], size=12, color=PC_LIGHT)

    # ── Slide 3: Experimental Setup ───────────────────────────────────────
    sl = add_slide(prs)
    set_bg(sl, PC_NAVY)
    add_rect(sl, 0, 0, 0.18, H_IN, PC_INDIGO)
    slide_tag(sl, "Experimental Setup")
    slide_title(sl, "8 Configurations Benchmarked")
    add_text(sl, 0.5, 1.25, 12, 0.3,
             ["2 feature representations  x  4 classification algorithms"],
             size=13, color=PC_LIGHT)
    pptx_table(sl, 0.5, 1.6, 6.1, 1.6, [
        ["Feature", "Dimensions", "Description"],
        ["Pixels",  "784-d",  "Flattened 28x28 normalized values"],
        ["HOG",     "324-d",  "9 orient. / 7x7px cell / 2x2 blocks / L2-Hys"],
    ], fs=10)
    pptx_table(sl, 0.5, 3.45, 6.1, 1.9, [
        ["Model", "Best Pixels Config", "Best HOG Config"],
        ["Logistic Regression", "C=0.1, lbfgs",   "C=0.1, lbfgs"],
        ["k-NN",               "k=7, distance",   "k=5, distance"],
        ["Random Forest",      "200 trees, d=20", "200 trees, d=inf"],
        ["Linear SVM",         "C=0.1, 2k iter",  "C=1.0, 2k iter"],
    ], fs=10)
    add_rect(sl, 6.9, 1.6, 6.1, 3.75, PC_SLATE)
    add_text(sl, 7.05, 1.7, 5.7, 0.3, ["EVALUATION METRICS"], size=8, color=PC_ACCENT, bold=True)
    add_text(sl, 7.05, 2.0, 5.7, 3.2, [
        "Accuracy (overall % correct on test set)",
        "Macro F1 (unweighted mean across 10 classes)",
        "Per-class Recall (sensitivity per category)",
        "Parameter Count:",
        "  - Weights for linear models",
        "  - MiB of training data for k-NN",
        "  - Node count for Random Forest",
        "Inference Latency (ms/image, median of 5 runs,",
        "  on 1,000-image batches)",
    ], size=11, color=PC_LIGHT)

    # ── Slide 4: Master Results ───────────────────────────────────────────
    sl = add_slide(prs)
    set_bg(sl, PC_NAVY)
    add_rect(sl, 0, 0, 0.18, H_IN, PC_INDIGO)
    slide_tag(sl, "Results")
    slide_title(sl, "Master Comparison: All 8 Configurations")
    add_text(sl, 0.5, 1.2, 12, 0.3,
             ["Ranked by test accuracy on 10,000 held-out images"],
             size=13, color=PC_LIGHT)
    sorted_r = sorted(results.items(), key=lambda x: -x[1]["test_accuracy"])
    tbl = [["#", "Model", "Features", "Test Acc", "Macro F1", "Params", "Latency (ms)"]]
    for i, ((m, f), r) in enumerate(sorted_r, 1):
        params = r["parameter_count"]
        ps = (f"{params:.1f} MiB" if m=="knn" else
              (f"{int(params):,} nd" if m=="rf" else f"{int(params):,}"))
        tbl.append([str(i), MODEL_LABELS[m], "Pix" if f=="pixels" else "HOG",
                    f"{r['test_accuracy']*100:.2f}%",
                    f"{r['test_macro_f1']*100:.2f}%", ps,
                    f"{r['inference_latency_ms_per_image']:.4f}"])
    pptx_table(sl, 0.5, 1.6, 12.5, 5.5, tbl, fs=9.5)

    # ── Slide 5: Key Findings ─────────────────────────────────────────────
    sl = add_slide(prs)
    set_bg(sl, PC_NAVY)
    add_rect(sl, 0, 0, 0.18, H_IN, PC_INDIGO)
    slide_tag(sl, "Key Findings")
    slide_title(sl, "5 Key Takeaways from the Comparison")
    findings = [
        ("1", "Random Forest dominates overall",
         "87.43% accuracy, +2.02pp over k-NN, +3.63pp over SVM. Non-linear boundaries from 200 trees are decisive."),
        ("2", "HOG helps linear models, hurts non-linear ones",
         "LogReg/SVM gain +0.23-0.43% with HOG (and 58.6% fewer params). RF/k-NN lose 3.08/1.57pp."),
        ("3", "k-NN memory/latency trade-off",
         "HOG cuts k-NN memory 58.7% (179->74 MiB) and latency 2.4x, at only 1.57pp accuracy cost."),
        ("4", "Linear classifiers are fastest",
         "LogReg+HOG: 0.0004 ms/img -- 685x faster than k-NN. Essential for real-time inventory tagging."),
        ("5", "Shirt (class 6) universally hardest",
         "Mean recall 57.7% across all 8 configs. No model exceeds 62.6%. A representation problem, not tuning."),
    ]
    for i, (num, heading, detail) in enumerate(findings):
        y = 1.42 + i * 1.05
        add_rect(sl, 0.5, y, 0.5, 0.8, PC_INDIGO)
        add_text(sl, 0.5, y + 0.15, 0.5, 0.5, [num], size=20, color=PC_WHITE,
                 bold=True, align=PP_ALIGN.CENTER)
        add_text(sl, 1.15, y,       11.5, 0.35, [heading], size=12, color=PC_ACCENT, bold=True)
        add_text(sl, 1.15, y + 0.38, 11.5, 0.65, [detail], size=10.5, color=PC_LIGHT)

    # ── Slide 6: Per-Class Recall ─────────────────────────────────────────
    sl = add_slide(prs)
    set_bg(sl, PC_NAVY)
    add_rect(sl, 0, 0, 0.18, H_IN, PC_INDIGO)
    slide_tag(sl, "Per-Class Analysis")
    slide_title(sl, "Per-Class Recall -- Best Model (RF + Pixels, 87.43%)")

    r_best = results[("rf", "pixels")]
    bar_data = sorted(zip(CLASS_NAMES, r_best["per_class_recall"]), key=lambda x: -x[1])
    rc_tbl = [["Class", "Recall", "Assessment"]]
    for cl, rec in bar_data:
        asmt = ("Excellent" if rec>=0.93 else "Good" if rec>=0.80 else "Fair" if rec>=0.70 else "Poor")
        rc_tbl.append([cl, f"{rec*100:.1f}%", asmt])
    pptx_table(sl, 0.5, 1.4, 5.5, 5.7, rc_tbl, fs=10)

    add_rect(sl, 6.3, 1.4, 6.7, 3.3, PC_SLATE)
    add_text(sl, 6.45, 1.5, 6.3, 0.3, ["CROSS-MODEL SHIRT RECALL"], size=8, color=PC_ACCENT, bold=True)
    shirt_lines = []
    for m_ in ["logreg","knn","rf","svm"]:
        rp = results[(m_,"pixels")]["per_class_recall"][6]
        rh = results[(m_,"hog")]["per_class_recall"][6]
        shirt_lines.append(f"{MODEL_LABELS[m_]}:  Pixels {rp*100:.1f}%  |  HOG {rh*100:.1f}%")
    mp = sum(results[(m_,"pixels")]["per_class_recall"][6] for m_ in ["logreg","knn","rf","svm"])/4
    mh = sum(results[(m_,"hog")]["per_class_recall"][6] for m_ in ["logreg","knn","rf","svm"])/4
    shirt_lines.append(f"Mean:  Pixels {mp*100:.1f}%  |  HOG {mh*100:.1f}%")
    add_text(sl, 6.45, 1.85, 6.3, 2.6, shirt_lines, size=11, color=PC_LIGHT)

    add_rect(sl, 6.3, 4.9, 6.7, 2.2, RGBColor(0x45, 0x0a, 0x0a))
    add_text(sl, 6.45, 5.0, 6.3, 0.3, ["THE SHIRT CEILING"], size=9, color=PC_RED, bold=True)
    add_text(sl, 6.45, 5.35, 6.3, 1.65,
             ["No model exceeds 62.6% Shirt recall.",
              "All hyperparameter configurations converge to the",
              "same ceiling -- a fundamental representation limit,",
              "not a tuning gap."],
             size=11, color=PC_LIGHT)

    # ── Slide 7: Error Analysis ───────────────────────────────────────────
    sl = add_slide(prs)
    set_bg(sl, PC_NAVY)
    add_rect(sl, 0, 0, 0.18, H_IN, PC_INDIGO)
    slide_tag(sl, "Error Analysis")
    slide_title(sl, "The Shirt Dilemma")
    add_text(sl, 0.5, 1.2, 12, 0.3,
             ["Why fixed descriptors fail for fine-grained upper-body classification"],
             size=13, color=PC_LIGHT)

    add_rect(sl, 0.5, 1.65, 5.9, 2.5, PC_SLATE)
    add_text(sl, 0.65, 1.75, 5.5, 0.3, ["CONFUSION PATTERN"], size=8, color=PC_ACCENT, bold=True)
    add_text(sl, 0.65, 2.05, 5.5, 2.0, [
        "Shirt -> T-shirt/top :  ~15-16% of Shirt test images",
        "Shirt -> Pullover    :  ~12-14% of Shirt test images",
        "Pullover -> Coat     :  ~8-10% of Pullover test images",
        "",
        "At 28x28 px, collar shapes, button rows, plackets",
        "and fabric weave become near-indistinguishable.",
    ], size=11, color=PC_LIGHT)

    rects = [
        ("Flattened Pixels",
         "784 independent positions -- no translation invariance. A collar shifted upward changes many coordinates."),
        ("HOG Descriptors",
         "Captures global edges but cannot learn which neckline edges are most diagnostic. Cell pooling removes fine spatial cues."),
        ("Linear Boundaries (LogReg/SVM)",
         "A hyperplane cannot create discriminative evidence absent in the representation. Tuning C doesn't resolve feature overlap."),
    ]
    for i, (heading, detail) in enumerate(rects):
        y = 4.4 + i * 0.98
        add_rect(sl, 0.5, y, 5.9, 0.88, PC_SLATE)
        add_text(sl, 0.65, y + 0.05, 5.5, 0.3, [heading], size=10, color=PC_ACCENT, bold=True)
        add_text(sl, 0.65, y + 0.4,  5.5, 0.45, [detail], size=10, color=PC_LIGHT)

    add_rect(sl, 6.7, 1.65, 6.3, 5.6, PC_SLATE)
    add_text(sl, 6.85, 1.75, 5.9, 0.3, ["VISUAL EVIDENCE (SVM ERROR ANALYSIS)"], size=8, color=PC_ACCENT, bold=True)
    add_text(sl, 6.85, 2.1, 5.9, 5.0, [
        "Four artifacts confirm the finding auditably:",
        "",
        "1. svm_pixels_confusion_matrix.png",
        "   Full 10-class normalized confusion matrix",
        "",
        "2. svm_pixels_upper_body_confusion.png",
        "   Focused 4x4 sub-matrix:",
        "   T-shirt / Pullover / Coat / Shirt",
        "",
        "3. svm_pixels_2_as_4_gallery.png",
        "   Pullover images misclassified as Coat",
        "",
        "4. svm_pixels_6_as_0_gallery.png",
        "   Shirt images misclassified as T-shirt/top",
    ], size=10.5, color=PC_LIGHT)

    # ── Slide 8: Efficiency ───────────────────────────────────────────────
    sl = add_slide(prs)
    set_bg(sl, PC_NAVY)
    add_rect(sl, 0, 0, 0.18, H_IN, PC_INDIGO)
    slide_tag(sl, "Efficiency Analysis")
    slide_title(sl, "Accuracy vs. Speed vs. Size")

    eff_tbl = [["Model", "Feat", "Accuracy", "Latency (ms)", "Parameters"]]
    for (m, f), r in sorted(results.items(), key=lambda x: -x[1]["test_accuracy"]):
        params = r["parameter_count"]
        ps = (f"{params:.1f} MiB" if m=="knn" else
              (f"{int(params):,} nd" if m=="rf" else f"{int(params):,}"))
        eff_tbl.append([MODEL_LABELS[m], "HOG" if f=="hog" else "Pix",
                        f"{r['test_accuracy']*100:.2f}%",
                        f"{r['inference_latency_ms_per_image']:.4f}", ps])
    pptx_table(sl, 0.5, 1.5, 7.5, 5.5, eff_tbl, fs=10)

    add_rect(sl, 8.3, 1.5, 4.8, 5.5, PC_SLATE)
    add_text(sl, 8.45, 1.6, 4.5, 0.3, ["DEPLOYMENT GUIDANCE"], size=8, color=PC_ACCENT, bold=True)
    add_text(sl, 8.45, 1.95, 4.5, 4.8, [
        "Maximum Accuracy:",
        "  RF + Pixels  --  87.43% at 0.026 ms/img",
        "",
        "Real-Time (<1ms budget):",
        "  LogReg + HOG  --  84.55% at 0.0004 ms",
        "  SVM + HOG  --  84.23% at 0.0006 ms",
        "",
        "Smallest Model (edge deployment):",
        "  LogReg / SVM + HOG  --  3,250 params",
        "  (58.6% fewer than pixel variants)",
        "",
        "k-NN Caveat:",
        "  85.41% accuracy but 179 MiB storage",
        "  + 685x slower than LogReg.",
        "  Not suitable for production scale.",
    ], size=11, color=PC_LIGHT)

    # ── Slide 9: Phase 2 Bridge ───────────────────────────────────────────
    sl = add_slide(prs)
    set_bg(sl, PC_NAVY)
    add_rect(sl, 0, 0, 0.18, H_IN, PC_INDIGO)
    slide_tag(sl, "Phase 2 Preview")
    slide_title(sl, "Why a CNN Is the Logical Next Step")

    add_rect(sl, 0.5, 1.45, 12.5, 0.9, RGBColor(0x45, 0x0a, 0x0a))
    add_text(sl, 0.65, 1.52, 12, 0.35,
             ["The Gap: 57.7% mean Shirt recall -- across all 8 Phase 1 configurations"],
             size=10, color=PC_RED, bold=True)
    add_text(sl, 0.65, 1.82, 12, 0.35,
             ["This is a representation failure, not a hyperparameter failure. "
              "No tuning of fixed features resolves overlapping garment silhouettes at 28x28 px."],
             size=9.5, color=PC_LIGHT)

    cnn_items = [
        ("Shared Detectors",
         "Convolutions share detectors across locations. A collar edge detector fires wherever the collar appears, not at a fixed coordinate."),
        ("Hierarchical Composition",
         "Stacked layers compose: edges -> textures -> collar shapes -> garment semantic concept. Impossible with a single fixed transform."),
        ("Translation Tolerance",
         "Pooling ensures small garment position shifts don't corrupt the feature vector. Directly fixes the pixel flattening problem."),
        ("End-to-End Learning",
         "The network discovers which local patterns are most diagnostic: button rows, collar shapes, plackets -- all fixed in HOG."),
    ]
    for i, (heading, detail) in enumerate(cnn_items):
        x = 0.5 + (i % 2) * 6.5
        y = 2.6 + (i // 2) * 2.0
        add_rect(sl, x, y, 6.1, 1.8, PC_SLATE)
        add_text(sl, x+0.15, y+0.1,  5.8, 0.4, [heading], size=12, color=PC_ACCENT, bold=True)
        add_text(sl, x+0.15, y+0.55, 5.8, 1.2, [detail], size=10.5, color=PC_LIGHT)

    add_rect(sl, 0.5, 6.6, 12.5, 0.65, RGBColor(0x04, 0x2f, 0x1e))
    add_text(sl, 0.65, 6.67, 12, 0.5,
             ["Success Criteria:  Shirt recall > 62.6%  |  Reduced Shirt/T-shirt/Pullover confusion  |  Overall accuracy > 87.43%"],
             size=10, color=PC_GREEN, bold=True)

    # ── Slide 10: Conclusion ──────────────────────────────────────────────
    sl = add_slide(prs)
    set_bg(sl, PC_NAVY)
    add_rect(sl, 0, 0, 0.18, H_IN, PC_INDIGO)
    slide_tag(sl, "Conclusion")
    slide_title(sl, "Phase 1 Summary")
    add_text(sl, 0.5, 1.2, 12, 0.3,
             ["A complete, reproducible AML baseline for 10-class clothing classification"],
             size=13, color=PC_LIGHT)

    pptx_table(sl, 0.5, 1.6, 7.5, 3.1, [
        ["Dimension", "Winner", "Value"],
        ["Best Accuracy",     "RF + Pixels",           "87.43%"],
        ["Best Macro F1",     "RF + Pixels",           "87.31%"],
        ["Fastest Inference", "LogReg + HOG",          "0.0004 ms/img"],
        ["Most Compact",      "LogReg / SVM + HOG",    "3,250 parameters"],
        ["Hardest Category",  "Shirt (class 6)",       "57.7% mean recall"],
    ], fs=11)

    add_rect(sl, 0.5, 4.9, 7.5, 2.3, PC_SLATE)
    add_text(sl, 0.65, 5.0, 7.1, 0.3, ["PHASE 1 DELIVERABLES"], size=8, color=PC_ACCENT, bold=True)
    add_text(sl, 0.65, 5.3, 7.1, 1.8, [
        "✓  8 standardized JSON result files in reports/results/",
        "✓  Master comparison notebook with charts",
        "✓  Error analysis report with visual evidence (4 figures)",
        "✓  SVM confusion matrices & misclassification galleries",
        "✓  Unit tests for all model modules",
    ], size=11, color=PC_LIGHT)

    add_rect(sl, 8.3, 1.6, 4.8, 5.6, RGBColor(0x04, 0x2f, 0x1e))
    add_text(sl, 8.45, 1.7, 4.5, 0.3, ["CONCLUSION"], size=9, color=PC_GREEN, bold=True)
    add_text(sl, 8.45, 2.05, 4.5, 5.0, [
        "Phase 1 establishes a rigorous",
        "baseline: all four models plateau",
        "at the same fundamental limit.",
        "",
        "Fine-grained upper-body garment",
        "categories (Shirt, T-shirt, Pullover,",
        "Coat) require spatial feature",
        "composition that fixed descriptors",
        "and linear boundaries cannot provide.",
        "",
        "Phase 2 (CNN) will directly test",
        "whether learned hierarchical spatial",
        "features close the identified gap.",
    ], size=11, color=PC_LIGHT)

    prs.save(str(out_path))
    print(f"[PPTX] Written -> {out_path}")
    return out_path


# ── MAIN ────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    print("Loading results...")
    results = load_results()
    print(f"  Loaded {len(results)} result files.")

    print("Building PDF report...")
    build_pdf(results)

    print("Building PPTX presentation...")
    build_pptx(results)

    print("\nDone!")
    print(f"  PDF  -> {OUT_DIR / 'phase1_report.pdf'}")
    print(f"  PPTX -> {OUT_DIR / 'phase1_slides.pptx'}")
